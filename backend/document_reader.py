import os

import pytesseract
from pdf2image import convert_from_path

POPPLER_PATH = (
    r"C:\Users\User\AppData\Local\Microsoft\WinGet\Packages"
    r"\oschwartz10612.Poppler_Microsoft.Winget.Source_8wekyb3d8bbwe"
    r"\poppler-25.07.0\Library\bin"
)
pytesseract.pytesseract.tesseract_cmd = (
    r"C:\Program Files\Tesseract-OCR\tesseract.exe"
)

from pypdf import PdfReader
from docx import Document
from pptx import Presentation
from openpyxl import load_workbook


def read_pdf(file_path: str):
    if not os.path.exists(file_path):
        raise FileNotFoundError("PDF file was not found.")

    reader = PdfReader(file_path)

    pages = []

    for page in reader.pages:
        text = page.extract_text()

        if text and text.strip():
            pages.append(text.strip())

    normal_text = "\n\n".join(pages)

    # If the PDF already contains text, use it.
    if normal_text.strip():
        return normal_text

    # If there is no text, use OCR for scanned PDF pages.
    try:
        images = convert_from_path(
    file_path,
    dpi=200,
    poppler_path=POPPLER_PATH
)

        ocr_pages = []

        for page_number, image in enumerate(
            images,
            start=1
        ):
            text = pytesseract.image_to_string(
                image
            )

            if text and text.strip():
                ocr_pages.append(
                    f"Page {page_number}:\n{text.strip()}"
                )

        return "\n\n".join(ocr_pages)

    except Exception as error:
        raise RuntimeError(
            "OCR could not read this PDF: "
            + str(error)
        )

def read_docx(file_path: str):
    if not os.path.exists(file_path):
        raise FileNotFoundError("Word document was not found.")

    document = Document(file_path)

    parts = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text)

    for table in document.tables:
        for row in table.rows:
            row_text = []

            for cell in row.cells:
                row_text.append(cell.text.strip())

            parts.append(" | ".join(row_text))

    return "\n".join(parts)


def read_pptx(file_path: str):
    if not os.path.exists(file_path):
        raise FileNotFoundError("PowerPoint file was not found.")

    presentation = Presentation(file_path)

    slides = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1
    ):
        slide_text = [
            f"Slide {slide_number}:"
        ]

        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text = shape.text.strip()

                if text:
                    slide_text.append(text)

        slides.append("\n".join(slide_text))

    return "\n\n".join(slides)


def read_xlsx(file_path: str):
    if not os.path.exists(file_path):
        raise FileNotFoundError("Excel file was not found.")

    workbook = load_workbook(
        file_path,
        data_only=True
    )

    sheets = []

    for worksheet in workbook.worksheets:

        sheets.append(
            f"Sheet: {worksheet.title}"
        )

        for row in worksheet.iter_rows(
            values_only=True
        ):
            values = []

            for value in row:
                if value is not None:
                    values.append(str(value))

            if values:
                sheets.append(
                    " | ".join(values)
                )

    return "\n".join(sheets)


def read_txt(file_path: str):
    if not os.path.exists(file_path):
        raise FileNotFoundError("Text file was not found.")

    with open(
        file_path,
        "r",
        encoding="utf-8",
        errors="ignore"
    ) as file:
        return file.read()


def read_document(file_path: str):
    """
    Read supported document types.
    """

    extension = os.path.splitext(
        file_path
    )[1].lower()

    if extension == ".pdf":
        return read_pdf(file_path)

    if extension == ".docx":
        return read_docx(file_path)

    if extension == ".pptx":
        return read_pptx(file_path)

    if extension == ".xlsx":
        return read_xlsx(file_path)

    if extension == ".txt":
        return read_txt(file_path)

    raise ValueError(
        "Unsupported document type: "
        + extension
    )


def get_document_info(file_path: str):
    """
    Return basic document information.
    """

    if not os.path.exists(file_path):
        raise FileNotFoundError(
            "Document was not found."
        )

    extension = os.path.splitext(
        file_path
    )[1].lower()

    return {
        "file": os.path.basename(file_path),
        "type": extension
    }